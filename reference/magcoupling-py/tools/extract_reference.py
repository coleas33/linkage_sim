"""Snapshot the workbook's cached values into tests/reference_values.json.

Usage:  python tools/extract_reference.py path/to/magnetic_coupling_torque_calculator.xlsx
Needs openpyxl. Run it again whenever the workbook is edited and recalculated,
then run the parity tests to see where the Python port and the workbook differ.
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

SHEETS = ["Calculator", "Calibration", "Metal design", "Materials", "Temperature design",
          "Shaft clamps", "Clamp screw sizes", "Gap sweep", "Pole sweep"]


def main(path):
    wb = load_workbook(path, data_only=True)
    ref = {}
    for name in SHEETS:
        ws = wb[name]
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and not isinstance(c.value, bool):
                    ref[f"{name}!{c.coordinate}"] = c.value
    out = Path(__file__).resolve().parents[1] / "tests" / "reference_values.json"
    out.write_text(json.dumps(ref, indent=0, sort_keys=True, default=str))
    print(f"wrote {len(ref)} cells to {out}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "magnetic_coupling_torque_calculator.xlsx")
